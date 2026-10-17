# -*- coding: utf-8 -*-
"""Verde Urbano — 10 slides para apresentação de extensão."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import nsmap
from lxml import etree

ROOT = Path(__file__).resolve().parent
SHOTS = ROOT / "shots"
OUT = ROOT / "Verde_Urbano_Apresentacao.pptx"

# LAYOUT_WIDE-ish: 13.333" x 7.5"
W, H = Inches(13.333), Inches(7.5)

FOREST = RGBColor(0x0F, 0x3D, 0x28)
GREEN = RGBColor(0x1A, 0x5C, 0x3A)
LEAF = RGBColor(0x2D, 0x8A, 0x58)
MOSS = RGBColor(0x97, 0xBC, 0x62)
GOLD = RGBColor(0xE8, 0xB8, 0x4A)
HEAT = RGBColor(0xC4, 0x4B, 0x2F)
INK = RGBColor(0x1A, 0x1F, 0x1C)
MUTED = RGBColor(0x5C, 0x6B, 0x62)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SOFT = RGBColor(0xF4, 0xF7, 0xF5)


def set_run(run, size, bold=False, color=INK, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_rect(slide, x, y, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    return sh


def add_round(slide, x, y, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    return sh


def add_oval(slide, x, y, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    return sh


def add_text(slide, text, x, y, w, h, size=16, bold=False, color=INK, font="Calibri", align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size, bold, color, font)
    # kill default margins
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    return box


def shot(name):
    p = SHOTS / name
    return p if p.exists() else None


def blank_slide(prs, bg=WHITE):
    s = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    add_rect(s, 0, 0, W, H, bg)
    return s


def main():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    # 1 CAPA
    s = blank_slide(prs, FOREST)
    add_oval(s, Inches(9.2), Inches(-1.5), Inches(6), Inches(6), GREEN)
    add_oval(s, Inches(-2), Inches(4.5), Inches(5), Inches(5), LEAF)
    uni = SHOTS / "uninassau-logo.png"
    if uni.exists():
        s.shapes.add_picture(str(uni), Inches(11.15), Inches(0.28), Inches(1.85), Inches(1.33))
    add_text(s, "UNINASSAU · TRABALHO DE EXTENSÃO · 2026", Inches(0.7), Inches(1.15), Inches(10), Inches(0.35),
             12, True, MOSS)
    add_text(s, "Verde Urbano", Inches(0.7), Inches(1.7), Inches(11), Inches(1.0),
             50, True, WHITE, "Cambria")
    add_text(s, "Mapear. Apoiar. Plantar.\nJuntos por uma Teresina mais verde.",
             Inches(0.7), Inches(2.9), Inches(9), Inches(1.0), 20, False, MOSS)
    add_text(s, "Yuri Medeiros Bandeira — 01693923\n"
                "Emanuel Da Silva Oliveira — 01679028\n"
                "Isaac Soares Almeida — 01694978\n"
                "Andressa Vitória dos Santos Gualter — 01701260",
             Inches(0.7), Inches(4.6), Inches(10), Inches(1.4),
             14, False, MOSS)
    add_text(s, "Teresina/PI  ·  vu-teresina.vercel.app",
             Inches(0.7), Inches(6.5), Inches(12), Inches(0.35),
             13, False, MOSS)

    # 2 PROBLEMA
    s = blank_slide(prs, WHITE)
    add_text(s, "O PROBLEMA", Inches(0.7), Inches(0.4), Inches(6), Inches(0.3), 12, True, LEAF)
    add_text(s, "Teresina precisa de sombra.", Inches(0.7), Inches(0.8), Inches(11.5), Inches(0.65),
             34, True, FOREST, "Cambria")
    add_text(
        s,
        "Uma das capitais mais quentes do Brasil — com pouca cobertura arbórea, praças degradadas e terrenos baldios sem canal estruturado de registro.",
        Inches(0.7), Inches(1.55), Inches(11.5), Inches(0.7), 16, False, MUTED,
    )
    cards = [
        ("14,6%", "copa média\nna área urbana", GREEN),
        ("46 °C", "temperatura\ndo chão na seca", HEAT),
        ("491", "quadras críticas\nsem sombra", FOREST),
    ]
    for i, (n, l, c) in enumerate(cards):
        x = Inches(0.7 + i * 4.1)
        add_round(s, x, Inches(2.6), Inches(3.8), Inches(3.5), SOFT)
        add_text(s, n, x + Inches(0.3), Inches(3.1), Inches(3.2), Inches(1.1), 46, True, c, "Cambria")
        add_text(s, l, x + Inches(0.3), Inches(4.5), Inches(3.2), Inches(1.0), 16, False, MUTED)
    add_text(s, "Fonte: ESA WorldCover + Landsat 8/9 · 1.884 hexágonos H3",
             Inches(0.7), Inches(6.85), Inches(11), Inches(0.3), 11, False, MUTED)

    # 3 OBJETIVOS
    s = blank_slide(prs, SOFT)
    add_text(s, "OBJETIVOS", Inches(0.7), Inches(0.4), Inches(6), Inches(0.3), 12, True, LEAF)
    add_text(s, "O que o MVP precisa validar", Inches(0.7), Inches(0.8), Inches(11), Inches(0.55),
             32, True, FOREST, "Cambria")
    goals = [
        ("Mapear", "Registrar locais degradados com foto, categoria e GPS."),
        ("Priorizar", "Mapa de prioridade com dados abertos de satélite."),
        ("Mobilizar", "Apoio coletivo e mutirões voluntários."),
        ("Orientar", "Guias SEMAM de espécies e competências dos órgãos."),
        ("Decidir", "Painel de gestão com indicadores e GeoJSON."),
        ("Sustentar", "Infraestrutura a custo zero após o semestre."),
    ]
    for i, (t, d) in enumerate(goals):
        col, row = i % 3, i // 3
        x = Inches(0.7 + col * 4.1)
        y = Inches(1.7 + row * 2.5)
        add_round(s, x, y, Inches(3.85), Inches(2.2), WHITE)
        add_oval(s, x + Inches(0.3), y + Inches(0.35), Inches(0.45), Inches(0.45), GREEN)
        add_text(s, str(i + 1), x + Inches(0.3), y + Inches(0.4), Inches(0.45), Inches(0.35),
                 14, True, WHITE, align=PP_ALIGN.CENTER)
        add_text(s, t, x + Inches(0.9), y + Inches(0.35), Inches(2.7), Inches(0.4),
                 20, True, FOREST, "Cambria")
        add_text(s, d, x + Inches(0.3), y + Inches(1.05), Inches(3.25), Inches(0.85),
                 14, False, MUTED)

    # 4 SOLUÇÃO
    s = blank_slide(prs, WHITE)
    add_text(s, "A SOLUÇÃO", Inches(0.7), Inches(0.4), Inches(6), Inches(0.3), 12, True, LEAF)
    add_text(s, "Uma plataforma, três atores", Inches(0.7), Inches(0.8), Inches(11), Inches(0.55),
             32, True, FOREST, "Cambria")
    add_text(
        s,
        "PWA que conecta cidadão, ONG de arborização e poder público para identificar, priorizar e transformar áreas sem vegetação.",
        Inches(0.7), Inches(1.45), Inches(11.5), Inches(0.55), 16, False, MUTED,
    )
    actors = [
        ("Cidadão", "Registra pontos, apoia demandas, entra em mutirões e consulta guias.", False),
        ("ONG", "Orienta espécies, organiza plantios e acompanha a revitalização.", True),
        ("Gestão", "Painel com mapa de inteligência, triagem e exportação GeoJSON.", False),
    ]
    for i, (t, d, hi) in enumerate(actors):
        x = Inches(0.7 + i * 4.1)
        add_round(s, x, Inches(2.35), Inches(3.85), Inches(3.7), GREEN if hi else SOFT)
        tc, dc = (WHITE, MOSS) if hi else (FOREST, MUTED)
        add_text(s, t, x + Inches(0.35), Inches(2.9), Inches(3.15), Inches(0.55),
                 26, True, tc, "Cambria")
        add_text(s, d, x + Inches(0.35), Inches(3.7), Inches(3.15), Inches(1.8),
                 16, False, dc)

    # 5 MAPA
    s = blank_slide(prs, FOREST)
    add_text(s, "DIFERENCIAL", Inches(0.55), Inches(0.3), Inches(5), Inches(0.28), 12, True, MOSS)
    add_text(s, "Onde plantar primeiro", Inches(0.55), Inches(0.65), Inches(8), Inches(0.5),
             28, True, WHITE, "Cambria")
    img = shot("mapear.png") or shot("01-mapear.png")
    if img:
        s.shapes.add_picture(str(img), Inches(0.45), Inches(1.3), Inches(8.5), Inches(5.6))
    add_round(s, Inches(9.2), Inches(1.3), Inches(3.65), Inches(5.6), GREEN)
    facts = [
        ("1.884", "hexágonos H3"),
        ("10 m", "resolução satélite"),
        ("3 lentes", "Prioridade · Copa · Calor"),
        ("ESA + USGS", "dados abertos reais"),
    ]
    for i, (n, l) in enumerate(facts):
        y = Inches(1.65 + i * 1.25)
        add_text(s, n, Inches(9.5), y, Inches(3.1), Inches(0.45), 22, True, GOLD, "Cambria")
        add_text(s, l, Inches(9.5), y + Inches(0.45), Inches(3.1), Inches(0.4), 13, False, MOSS)

    # 6 LANDING
    s = blank_slide(prs, SOFT)
    add_text(s, "PRODUTO", Inches(0.55), Inches(0.28), Inches(5), Inches(0.28), 12, True, LEAF)
    add_text(s, "A cara do app", Inches(0.55), Inches(0.6), Inches(8), Inches(0.45),
             28, True, FOREST, "Cambria")
    img = shot("landing.png") or shot("02-landing.png")
    if img:
        s.shapes.add_picture(str(img), Inches(0.45), Inches(1.2), Inches(12.4), Inches(5.8))

    # 7 FEED + MUTIRÕES
    s = blank_slide(prs, WHITE)
    add_text(s, "CIDADÃO", Inches(0.5), Inches(0.25), Inches(5), Inches(0.28), 12, True, LEAF)
    add_text(s, "Do pedido ao plantio", Inches(0.5), Inches(0.55), Inches(12), Inches(0.45),
             28, True, FOREST, "Cambria")
    feed = shot("feed.png")
    mut = shot("mutiroes.png")
    if feed:
        s.shapes.add_picture(str(feed), Inches(0.35), Inches(1.15), Inches(6.25), Inches(5.85))
    if mut:
        s.shapes.add_picture(str(mut), Inches(6.8), Inches(1.15), Inches(6.15), Inches(5.85))

    # 8 GUIAS
    s = blank_slide(prs, SOFT)
    add_text(s, "ORIENTAÇÃO", Inches(0.55), Inches(0.28), Inches(5), Inches(0.28), 12, True, LEAF)
    add_text(s, "Guias SEMAM e espécies nativas", Inches(0.55), Inches(0.6), Inches(10), Inches(0.45),
             28, True, FOREST, "Cambria")
    guias = shot("guias.png")
    if guias:
        s.shapes.add_picture(str(guias), Inches(0.45), Inches(1.25), Inches(8.35), Inches(5.7))
    add_round(s, Inches(9.05), Inches(1.25), Inches(3.8), Inches(5.7), WHITE)
    add_text(s, "O que o cidadão encontra", Inches(9.35), Inches(1.6), Inches(3.2), Inches(0.7),
             16, True, FOREST, "Cambria")
    for i, tip in enumerate([
        "Espécies nativas / adaptadas",
        "Porte e tipo de raiz",
        "Viveiros municipais",
        "Órgãos competentes",
        "Fontes oficiais SEMAM",
    ]):
        add_text(s, "→  " + tip, Inches(9.35), Inches(2.6 + i * 0.7), Inches(3.2), Inches(0.5),
                 14, False, MUTED)

    # 9 RESULTADOS
    s = blank_slide(prs, WHITE)
    add_text(s, "RESULTADOS", Inches(0.7), Inches(0.4), Inches(6), Inches(0.3), 12, True, LEAF)
    add_text(s, "O que entregamos", Inches(0.7), Inches(0.8), Inches(11), Inches(0.55),
             32, True, FOREST, "Cambria")
    results = [
        ("Mapa verde inédito", "Prioridade de arborização de Teresina com dados abertos e reprodutíveis."),
        ("App publicado", "PWA em produção — vu-teresina.vercel.app — instalável no celular."),
        ("Custo zero", "Next.js, MapLibre, Vercel e dados abertos — sustentável após o semestre."),
        ("Hipótese validada", "Dados + participação cidadã orientam política de arborização urbana."),
    ]
    for i, (t, d) in enumerate(results):
        y = Inches(1.6 + i * 1.3)
        add_round(s, Inches(0.7), y, Inches(11.9), Inches(1.15), SOFT)
        add_round(s, Inches(0.95), y + Inches(0.3), Inches(0.55), Inches(0.55), GREEN)
        add_text(s, str(i + 1), Inches(0.95), y + Inches(0.38), Inches(0.55), Inches(0.4),
                 16, True, WHITE, align=PP_ALIGN.CENTER)
        add_text(s, t, Inches(1.75), y + Inches(0.2), Inches(10.5), Inches(0.35),
                 18, True, FOREST, "Cambria")
        add_text(s, d, Inches(1.75), y + Inches(0.58), Inches(10.5), Inches(0.4),
                 14, False, MUTED)

    # 10 FECHO
    s = blank_slide(prs, FOREST)
    add_oval(s, Inches(8.5), Inches(-2), Inches(7), Inches(7), GREEN)
    if uni.exists():
        s.shapes.add_picture(str(uni), Inches(11.15), Inches(0.28), Inches(1.85), Inches(1.33))
    add_text(s, "OBRIGADO", Inches(0.8), Inches(1.5), Inches(11), Inches(0.4), 14, True, MOSS)
    add_text(s, "Mapear. Apoiar. Plantar.", Inches(0.8), Inches(2.0), Inches(11.5), Inches(0.85),
             40, True, WHITE, "Cambria")
    add_text(s, "Juntos por uma Teresina mais verde.", Inches(0.8), Inches(2.95), Inches(11), Inches(0.5),
             22, False, MOSS)
    add_round(s, Inches(0.8), Inches(3.85), Inches(5.6), Inches(1.35), GREEN)
    add_text(s, "Acesse o app", Inches(1.1), Inches(4.0), Inches(5), Inches(0.35), 13, False, MOSS)
    add_text(s, "vu-teresina.vercel.app", Inches(1.1), Inches(4.45), Inches(5), Inches(0.5),
             22, True, WHITE)
    add_text(s, "Yuri Medeiros Bandeira — 01693923  ·  Emanuel Da Silva Oliveira — 01679028\n"
                "Isaac Soares Almeida — 01694978  ·  Andressa Vitória dos Santos Gualter — 01701260\n"
                "UNINASSAU · Extensão 2026 · Teresina/PI",
             Inches(0.8), Inches(5.55), Inches(11.5), Inches(1.1), 12, False, MOSS)

    prs.save(OUT)
    print("OK", OUT)


if __name__ == "__main__":
    main()
